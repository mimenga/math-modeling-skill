# -*- coding: utf-8 -*-
"""生成国赛论文 Word 模板 (python-docx)

对齐官方 2026 修订稿：
- 电子版第一页为摘要专用页（不自制封面、不放姓名——匿名要求第六条）；
- 正文不要目录（第四条）；
- 三线表（团队内部约定，与 05_论文格式规范.md 一致）；
- 公式全文连续编号示例。
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- 页面设置 A4 ----------
for section in doc.sections:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# ---------- 页眉：队伍编号（不放姓名/学校，匿名要求） ----------
header = doc.sections[0].header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run('队伍编号：XXXXX')
hr.font.size = Pt(9)
hr.font.name = 'Times New Roman'
hr.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')   # 页眉中文需显式 eastAsia 字体

# ---------- 页脚：页码从摘要页起居中 ----------
footer = doc.sections[0].footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
fp._p.append(fld)

# ---------- 默认字体：宋体小四 ----------
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_heading_cn(text, level=1):
    """中文标题：黑体"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, align=None, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def style_three_line(table):
    """三线表：顶线/底线 1.5 磅，表头下线 0.75 磅，无竖线。"""
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '12')          # 单位 1/8 磅，12 = 1.5 磅
        el.set(qn('w:color'), '000000')
        borders.append(el)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'none')
        borders.append(el)
    tblPr.insert_element_before(borders, 'w:shd', 'w:tblLayout',
                                'w:tblCellMar', 'w:tblLook',
                                'w:tblCaption', 'w:tblDescription')
    for cell in table.rows[0].cells:      # 表头下线 0.75 磅
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        el = OxmlElement('w:bottom')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')           # 6 = 0.75 磅
        el.set(qn('w:color'), '000000')
        tcBorders.append(el)
        tcPr.insert_element_before(tcBorders, 'w:shd', 'w:noWrap', 'w:tcMar',
                                   'w:textDirection', 'w:tcFitText',
                                   'w:vAlign', 'w:hideMark')

def make_table(headers, rows, col_widths=None):
    """创建三线表并填充内容。行不跨页断裂、跨页自动重复表头。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            table.rows[i].cells[j].text = val
    style_three_line(table)
    # 防跨页断裂：每行 cantSplit（行内不折页）；首行 tblHeader（跨页重复表头）
    for k, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
        if k == 0:
            th = OxmlElement('w:tblHeader')
            trPr.append(th)
    return table

def add_formula(number, desc='【公式：用 MathType 输入，变量斜体】'):
    """带右侧编号的独立公式占位：居中制表位放公式，右制表位放编号。"""
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(f'\t{desc}\t({number})')
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return p

def add_blank():
    doc.add_paragraph()

# ============================================================
# 摘要专用页（电子版第一页；承诺书/编号专用页由组委会提供，勿自制）
# ============================================================
add_para('论文题目（居中，一句话概括核心工作）', bold=True, align='center', size=16)
add_para('摘  要', bold=True, align='center', size=16)
add_blank()
add_para('本文针对【问题背景一句话】，建立了【核心模型名】模型。', size=12)
add_para('针对问题一，本文采用【方法】方法，建立【模型】模型，得到【结果】（关键数字：___）。', size=12)
add_para('针对问题二，本文构建了【模型】模型，求解得【具体数值结果，如最优时点/最优方案】。', size=12)
add_para('针对问题三，在问题二基础上引入【因素】，建立【模型】，得到【结果】。', size=12)
add_para('针对问题四，本文提出【方法】判定方法，阈值设为___，灵敏度___，特异度___。', size=12)
add_para('最后，对模型进行了灵敏度分析（关键参数 ±5%、±10% 变化，结果波动___），表明模型具有较好的鲁棒性。', size=12)
add_para('【提示：① 创新模型名/方法名用粗体标出；② 题目要求交付结果文件时写明"计算结果存于 result1.xlsx 等"；③ 每问段末落 2~3 个关键数字，不要只给一个孤立数字。】', size=10)
add_blank()
add_para('关键词：___；___；___；___；___', size=12)
doc.add_page_break()

# ============================================================
# 正文（官方规定正文不要目录，直接从问题重述开始）
# ============================================================
add_heading_cn('一、问题重述与分析', 1)
add_heading_cn('1.1 问题背景', 2)
add_para('【用自己的话描述问题背景，不要复制粘贴题目原文，半页以内】')
add_heading_cn('1.2 问题重述', 2)
for i in range(1, 5):
    add_para(f'问题{i}：【一句话概括该小问要求】')
add_heading_cn('1.3 问题分析', 2)
add_para('【给出整体解题思路，说明各问之间的联系，建议画全文建模思路总流程图并插入此处（通常为图1）】')

add_heading_cn('二、模型假设与符号说明', 1)
add_heading_cn('2.1 模型假设', 2)
for i in range(1, 4):
    add_para(f'假设{i}：【假设内容】。理由：【为什么这样假设，最好给出定量依据或文献支撑】')

add_heading_cn('2.2 符号说明', 2)
make_table(
    ['符号', '含义', '单位'],
    [
        ('$x_i$', '【含义】', '【单位】'),
        ('$y_j$', '【含义】', '【单位】'),
        ('$\\alpha$', '【含义】', '【单位】'),
        ('$\\beta$', '【含义】', '【单位】'),
    ],
)

add_heading_cn('三、数据预处理', 1)
add_heading_cn('3.1 数据来源与概况', 2)
add_para('【样本量、变量数、数据类型、时间范围】')
add_heading_cn('3.2 缺失值与异常值处理', 2)
add_para('【处理方法和理由：删除/插补/3σ/箱线图。重要：说明"为什么这样做"，也要说明"为什么不这样做"（如保留基线的理由）】')
add_heading_cn('3.3 特征工程与数据洞察', 2)
add_para('【哑变量、衍生变量（如 BMI）、标准化、PCA 等；以 1~2 条业务洞察收尾，不止罗列清洗操作】')

add_heading_cn('四、问题一的模型建立与求解', 1)
add_heading_cn('4.1 模型建立', 2)
add_para('【建模思路：先一句话引出公式，公式后逐符号解释并给出参数取值。公式编号从(1)起全文连续，不按章节编号；后文引用一律写"式(1)"】')
add_formula(1)
add_para('式中：【符号1】为【含义，取值___】；【符号2】为【含义】。', size=12)
add_formula(2, '【判别条件/约束条件也可编号，便于"由式(2)"引用】')
add_heading_cn('4.2 模型求解', 2)
add_para('【求解方法、工具与版本（如 YALMIP+Gurobi，MIPGap=1e-4；MATLAB R2023a）、算法流程或伪代码（粗体 STEP 1/2/3，每步"做什么+为什么"），完整代码放附录】')
add_heading_cn('4.3 结果与分析', 2)
make_table(
    ['方案/对象', '关键指标', 'RMSE', 'R²'],
    [
        ('【方案A】', '___', '___', '___'),
        ('【方案B】', '___', '___', '___'),
    ],
)
add_para('表：结果表建议带 RMSE/R²/相对误差列。图注可携带情景参数（如"图2 t=412.48 s 时位置图"）；统计表配显著性星号+表底注释行；大表写"仅展示部分，完整数据见附录"。之后配 2~3 句解读：为什么是这个结果、是否合理。', size=10)
add_para('【排版铁律：表题在表上方、图题在图下方，题注段与表/图必须同页（题注段落属性设"与下段同页"）；表格已设行不跨页断裂+跨页重复表头】', size=10)

add_heading_cn('五、问题二的模型建立与求解', 1)
add_heading_cn('5.1 模型建立', 2)
add_para('【优化三段式：（1）决策变量：…（2）目标函数：…（式号）（3）约束条件：逐条①②③，每条公式后跟一句物理/经济意义，最后汇总为 max Z  s.t. ①~⑥（式号）】')
add_heading_cn('5.2 模型求解', 2)
add_para('【求解算法 + 结果；迭代/搜索类算法给出"迭代过程表（次数|取值|目标值|收敛方向）+ 收敛图"】')
add_heading_cn('5.3 结果与分析', 2)
add_para('【最优方案明确列出 + 图表 + 解读；结果反常时主动写一段归因（现象→约束/假设→"符合逻辑"）】')

add_heading_cn('六、问题三的模型建立与求解', 1)
add_heading_cn('6.1 模型建立', 2)
add_para('【多因素升级，说明与问题二模型的区别】')
add_heading_cn('6.2 模型求解', 2)
add_heading_cn('6.3 结果与分析', 2)

add_heading_cn('七、问题四的模型建立与求解', 1)
add_heading_cn('7.1 模型建立', 2)
add_heading_cn('7.2 模型求解', 2)
add_heading_cn('7.3 结果与分析', 2)

add_heading_cn('八、模型的评价、改进与推广', 1)
add_heading_cn('8.1 灵敏度分析', 2)
add_para('【选 2~3 个关键参数：定扰动方式（±5%/分布/蒙特卡洛/对照实验/极值情景）→表格呈现（参数|基准|−5%|+5%|变化率）→结论三句式：定区间→量化变化→横向排序+"设计时应重点关注…"】')
add_heading_cn('8.2 模型的优点', 2)
add_para('【优点要具体，与模型机制绑定，不空夸】')
add_heading_cn('8.3 模型的缺点与改进方向', 2)
add_para('【缺点要具体、可量化（如"计算一次需 30 分钟以上""与理论最优解差距不超过 0.5%"），显诚实但避免致命伤；改进方向与缺点一一对应】')
add_heading_cn('8.4 模型的推广', 2)
add_para('【可用于哪些类似场景，落到具体行业，一句话】')

add_heading_cn('AI 工具使用声明', 1)
add_para('本参赛队在竞赛过程中使用了 AI 工具，主要用于【语言润色、代码调试、资料查询】等辅助工作，详细使用情况见支撑材料。')
add_para('【2026 新规硬性要求：本声明必须在参考文献之前。若未使用 AI，整段替换为：本参赛队在竞赛过程中未使用任何 AI 工具。】')
add_para('【使用 AI 的队还需另附支撑材料 PDF：AI工具使用详情.pdf，含工具名称与版本、使用环节、提示方式、人工核验情况。】')

add_heading_cn('参考文献', 1)
add_para('[1] 作者. 题目[J]. 期刊名, 年, 卷(期): 页码.')
add_para('[2] 作者. 书名[M]. 出版地: 出版社, 年份.')
add_para('[3] 作者. 文章题目[C]//会议论文集. 出版地: 出版社, 年份: 页码.')

add_heading_cn('附录', 1)
add_para('【附表 1：支撑材料三列表（文件名|类型|简介）。官方第五条：附录必须含支撑材料文件列表+全部完整可运行源程序，否则可能取消评奖资格】')
add_para('【代码放这里，不放正文。按问题命名+中文功能注释（如 Q1_Problem_1.m：问题一求解主程序）】')

_out_path = Path(__file__).resolve().parent / '国赛论文模板_Word.docx'
_out_path.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(_out_path))
print(f'Word 模板生成成功: {_out_path}')
